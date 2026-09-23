import json
import os
import secrets
import sqlite3
import uuid
from datetime import datetime, timezone
from io import BytesIO

from flask import Flask, g, jsonify, request, send_file, session, send_from_directory

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("DB_PATH", os.path.join(BASE_DIR, "data", "bidvest_esg.db"))
UPLOAD_DIR = os.environ.get("UPLOAD_DIR", os.path.join(BASE_DIR, "data", "uploads"))
CHECKLIST_JSON = os.path.join(BASE_DIR, "checklist_items.json")

TEAM_PASSWORD = os.environ.get("APP_PASSWORD_TEAM", "bidvest-esg-team")
ADMIN_PASSWORD = os.environ.get("APP_PASSWORD_ADMIN", "bidvest-esg-admin")
MAX_UPLOAD_MB = 20

os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)

app = Flask(__name__, static_folder="static", static_url_path="/static")
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024

DEFAULT_OWNERS = [
    "ESG Owner (Group liaison)",
    "Site Manager",
    "HR",
    "Procurement",
    "Finance",
    "Health & Safety",
    "Learning Academy",
]

STATUSES = ["Not started", "In progress", "Collected", "Verified", "N/A"]


# ---------------------------------------------------------------- database

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA journal_mode=WAL")
        g.db.execute("PRAGMA busy_timeout=5000")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA journal_mode=WAL")
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS checklist_items (
            id TEXT PRIMARY KEY,
            pillar TEXT NOT NULL,
            category TEXT NOT NULL,
            data_point TEXT NOT NULL,
            unit TEXT,
            frequency TEXT,
            typical_source TEXT,
            framework_ref TEXT,
            sort_order INTEGER
        );

        CREATE TABLE IF NOT EXISTS capture_records (
            item_id TEXT PRIMARY KEY REFERENCES checklist_items(id),
            status TEXT NOT NULL DEFAULT 'Not started',
            latest_value TEXT,
            period TEXT,
            notes TEXT,
            owner TEXT,
            updated_by TEXT,
            updated_at TEXT
        );

        CREATE TABLE IF NOT EXISTS capture_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_id TEXT NOT NULL,
            status TEXT,
            latest_value TEXT,
            period TEXT,
            notes TEXT,
            owner TEXT,
            updated_by TEXT,
            updated_at TEXT
        );

        CREATE TABLE IF NOT EXISTS evidence_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_id TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            original_name TEXT NOT NULL,
            filesize INTEGER,
            uploaded_by TEXT,
            uploaded_at TEXT
        );

        CREATE TABLE IF NOT EXISTS owners (
            name TEXT PRIMARY KEY
        );
        """
    )
    db.commit()

    # Seed checklist items idempotently. Multiple gunicorn workers can call
    # init_db() concurrently on first boot (same fresh sqlite file) — use
    # INSERT OR IGNORE everywhere here so a race never raises IntegrityError.
    with open(CHECKLIST_JSON) as f:
        items = json.load(f)
    for idx, item in enumerate(items):
        db.execute(
            """INSERT OR IGNORE INTO checklist_items
               (id, pillar, category, data_point, unit, frequency, typical_source, framework_ref, sort_order)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                item["id"], item["pillar"], item["category"], item["data_point"],
                item.get("unit"), item.get("frequency"), item.get("typical_source"),
                item.get("framework_ref"), idx,
            ),
        )
        db.execute(
            """INSERT OR IGNORE INTO capture_records (item_id, status, updated_at)
               VALUES (?, 'Not started', ?)""",
            (item["id"], datetime.now(timezone.utc).isoformat()),
        )
    db.commit()

    cur = db.execute("SELECT COUNT(*) AS c FROM owners")
    if cur.fetchone()["c"] == 0:
        for name in DEFAULT_OWNERS:
            db.execute("INSERT OR IGNORE INTO owners (name) VALUES (?)", (name,))
        db.commit()

    db.close()


init_db()


# ---------------------------------------------------------------- auth

def require_login(admin_only=False):
    role = session.get("role")
    if not role:
        return jsonify({"error": "Not logged in"}), 401
    if admin_only and role != "admin":
        return jsonify({"error": "Admin access required"}), 403
    return None


@app.route("/api/login", methods=["POST"])
def login():
    data = request.get_json(force=True, silent=True) or {}
    password = data.get("password", "")
    name = (data.get("name") or "").strip()
    if password == ADMIN_PASSWORD:
        session["role"] = "admin"
        session["name"] = name or "ESG Owner"
        return jsonify({"role": "admin", "name": session["name"]})
    if password == TEAM_PASSWORD:
        session["role"] = "team"
        session["name"] = name or "Team member"
        return jsonify({"role": "team", "name": session["name"]})
    return jsonify({"error": "Incorrect password"}), 401


@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/session")
def whoami():
    if not session.get("role"):
        return jsonify({"loggedIn": False})
    return jsonify({"loggedIn": True, "role": session["role"], "name": session.get("name")})


# ---------------------------------------------------------------- checklist

@app.route("/api/checklist")
def checklist():
    err = require_login()
    if err:
        return err
    db = get_db()
    rows = db.execute(
        """SELECT ci.*, cr.status, cr.latest_value, cr.period, cr.notes, cr.owner,
                  cr.updated_by, cr.updated_at,
                  (SELECT COUNT(*) FROM evidence_files ef WHERE ef.item_id = ci.id) AS evidence_count
           FROM checklist_items ci
           LEFT JOIN capture_records cr ON cr.item_id = ci.id
           ORDER BY ci.sort_order"""
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/owners", methods=["GET", "POST"])
def owners():
    db = get_db()
    if request.method == "GET":
        err = require_login()
        if err:
            return err
        rows = db.execute("SELECT name FROM owners ORDER BY name").fetchall()
        return jsonify([r["name"] for r in rows])
    err = require_login(admin_only=True)
    if err:
        return err
    data = request.get_json(force=True, silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Name required"}), 400
    db.execute("INSERT OR IGNORE INTO owners (name) VALUES (?)", (name,))
    db.commit()
    rows = db.execute("SELECT name FROM owners ORDER BY name").fetchall()
    return jsonify([r["name"] for r in rows])


@app.route("/api/capture", methods=["POST"])
def capture():
    err = require_login()
    if err:
        return err
    data = request.get_json(force=True, silent=True) or {}
    item_id = data.get("item_id")
    if not item_id:
        return jsonify({"error": "item_id required"}), 400
    status = data.get("status")
    if status not in STATUSES:
        return jsonify({"error": "Invalid status"}), 400

    db = get_db()
    exists = db.execute("SELECT 1 FROM checklist_items WHERE id = ?", (item_id,)).fetchone()
    if not exists:
        return jsonify({"error": "Unknown item_id"}), 404

    now = datetime.now(timezone.utc).isoformat()
    updated_by = session.get("name", "Unknown")
    latest_value = data.get("latest_value")
    period = data.get("period")
    notes = data.get("notes")
    owner = data.get("owner")

    db.execute(
        """INSERT INTO capture_records (item_id, status, latest_value, period, notes, owner, updated_by, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)
           ON CONFLICT(item_id) DO UPDATE SET
             status=excluded.status, latest_value=excluded.latest_value, period=excluded.period,
             notes=excluded.notes, owner=excluded.owner, updated_by=excluded.updated_by, updated_at=excluded.updated_at""",
        (item_id, status, latest_value, period, notes, owner, updated_by, now),
    )
    db.execute(
        """INSERT INTO capture_history (item_id, status, latest_value, period, notes, owner, updated_by, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (item_id, status, latest_value, period, notes, owner, updated_by, now),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/history/<item_id>")
def history(item_id):
    err = require_login()
    if err:
        return err
    db = get_db()
    rows = db.execute(
        "SELECT * FROM capture_history WHERE item_id = ? ORDER BY id DESC LIMIT 50", (item_id,)
    ).fetchall()
    return jsonify([dict(r) for r in rows])


# ---------------------------------------------------------------- evidence

ALLOWED_EXTENSIONS = {
    "pdf", "png", "jpg", "jpeg", "heic", "gif", "webp",
    "xlsx", "xls", "csv", "docx", "doc", "txt", "eml", "msg",
}


@app.route("/api/evidence/<item_id>")
def list_evidence(item_id):
    err = require_login()
    if err:
        return err
    db = get_db()
    rows = db.execute(
        "SELECT id, original_name, filesize, uploaded_by, uploaded_at FROM evidence_files WHERE item_id = ? ORDER BY id DESC",
        (item_id,),
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/upload", methods=["POST"])
def upload():
    err = require_login()
    if err:
        return err
    item_id = request.form.get("item_id")
    if not item_id:
        return jsonify({"error": "item_id required"}), 400
    db = get_db()
    exists = db.execute("SELECT 1 FROM checklist_items WHERE id = ?", (item_id,)).fetchone()
    if not exists:
        return jsonify({"error": "Unknown item_id"}), 404

    if "file" not in request.files:
        return jsonify({"error": "No file"}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "No file selected"}), 400
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        return jsonify({"error": f"File type .{ext} not allowed"}), 400

    stored_name = f"{item_id}_{uuid.uuid4().hex}.{ext}"
    dest_dir = os.path.join(UPLOAD_DIR, item_id)
    os.makedirs(dest_dir, exist_ok=True)
    dest_path = os.path.join(dest_dir, stored_name)
    file.save(dest_path)
    filesize = os.path.getsize(dest_path)

    now = datetime.now(timezone.utc).isoformat()
    uploaded_by = session.get("name", "Unknown")
    cur = db.execute(
        """INSERT INTO evidence_files (item_id, stored_name, original_name, filesize, uploaded_by, uploaded_at)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (item_id, stored_name, file.filename, filesize, uploaded_by, now),
    )
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/download/<int:file_id>")
def download(file_id):
    err = require_login()
    if err:
        return err
    db = get_db()
    row = db.execute("SELECT * FROM evidence_files WHERE id = ?", (file_id,)).fetchone()
    if not row:
        return jsonify({"error": "Not found"}), 404
    path = os.path.join(UPLOAD_DIR, row["item_id"], row["stored_name"])
    if not os.path.exists(path):
        return jsonify({"error": "File missing on disk"}), 410
    return send_file(path, as_attachment=True, download_name=row["original_name"])


@app.route("/api/evidence/<int:file_id>", methods=["DELETE"])
def delete_evidence(file_id):
    err = require_login(admin_only=True)
    if err:
        return err
    db = get_db()
    row = db.execute("SELECT * FROM evidence_files WHERE id = ?", (file_id,)).fetchone()
    if not row:
        return jsonify({"error": "Not found"}), 404
    path = os.path.join(UPLOAD_DIR, row["item_id"], row["stored_name"])
    if os.path.exists(path):
        os.remove(path)
    db.execute("DELETE FROM evidence_files WHERE id = ?", (file_id,))
    db.commit()
    return jsonify({"ok": True})


# ---------------------------------------------------------------- summary / export

def compute_summary(db):
    rows = db.execute(
        """SELECT ci.pillar, ci.category, ci.frequency, cr.status
           FROM checklist_items ci LEFT JOIN capture_records cr ON cr.item_id = ci.id"""
    ).fetchall()

    by_pillar = {}
    by_frequency = {}
    for r in rows:
        p = by_pillar.setdefault(r["pillar"], {s: 0 for s in STATUSES})
        p[r["status"] or "Not started"] += 1
        f = by_frequency.setdefault(r["frequency"] or "Unspecified", {"total": 0, "not_collected": 0})
        f["total"] += 1
        if r["status"] not in ("Collected", "Verified"):
            f["not_collected"] += 1

    pillar_summary = []
    total = {s: 0 for s in STATUSES}
    total_points = 0
    for pillar, counts in by_pillar.items():
        points = sum(counts.values())
        total_points += points
        denom = points - counts["N/A"]
        pct = round(100 * (counts["Collected"] + counts["Verified"]) / denom, 1) if denom else 0.0
        pillar_summary.append({"pillar": pillar, "data_points": points, **counts, "pct_complete": pct})
        for s in STATUSES:
            total[s] += counts[s]

    denom_total = total_points - total["N/A"]
    total_pct = round(100 * (total["Collected"] + total["Verified"]) / denom_total, 1) if denom_total else 0.0

    frequency_summary = [
        {"frequency": f, "data_points": v["total"], "not_yet_collected": v["not_collected"]}
        for f, v in by_frequency.items()
    ]

    evidence_row = db.execute(
        """SELECT COUNT(DISTINCT ci.id) AS with_evidence
           FROM checklist_items ci JOIN evidence_files ef ON ef.item_id = ci.id"""
    ).fetchone()

    return {
        "pillars": pillar_summary,
        "total": {"data_points": total_points, **total, "pct_complete": total_pct},
        "frequency": frequency_summary,
        "items_with_evidence": evidence_row["with_evidence"],
        "total_items": total_points,
    }


@app.route("/api/summary")
def summary():
    err = require_login()
    if err:
        return err
    return jsonify(compute_summary(get_db()))


@app.route("/api/export/xlsx")
def export_xlsx():
    err = require_login()
    if err:
        return err
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    db = get_db()
    rows = db.execute(
        """SELECT ci.*, cr.status, cr.latest_value, cr.period, cr.notes, cr.owner
           FROM checklist_items ci LEFT JOIN capture_records cr ON cr.item_id = ci.id
           ORDER BY ci.sort_order"""
    ).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ESG Checklist"
    headers = ["ID", "Pillar", "Category", "Data point", "Unit", "Frequency", "Typical source",
               "Framework / reference", "Owner", "Status", "Latest value", "Period", "Notes / evidence", "Evidence files"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")

    for r in rows:
        ev_count = db.execute("SELECT COUNT(*) c FROM evidence_files WHERE item_id=?", (r["id"],)).fetchone()["c"]
        ws.append([
            r["id"], r["pillar"], r["category"], r["data_point"], r["unit"], r["frequency"],
            r["typical_source"], r["framework_ref"], r["owner"], r["status"], r["latest_value"],
            r["period"], r["notes"], ev_count,
        ])
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 50)

    summary_data = compute_summary(db)
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Pillar", "Data points"] + STATUSES + ["% complete"])
    for p in summary_data["pillars"]:
        ws2.append([p["pillar"], p["data_points"]] + [p[s] for s in STATUSES] + [p["pct_complete"]])
    t = summary_data["total"]
    ws2.append(["Total", t["data_points"]] + [t[s] for s in STATUSES] + [t["pct_complete"]])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Bidvest_ESG_Checklist_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=fname,
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------- static / health

@app.route("/health")
def health():
    return jsonify({"status": "ok"}), 200


@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
